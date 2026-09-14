# Adapter: ExcelDocumentReader.
# Posjeduje: cita .xlsx/.xls, kesiranje, DocumentEvidence output.
# Implementira: ports/document_reader.py DocumentReader.
# Ne zna za: SQLite, PySide6, Docling, contract.
"""ExcelDocumentReader — adapter za .xlsx/.xls fajlove.

Generalizacija logike iz legacy core/documents/excel_document.py:
- otvara .xlsx sa openpyxl (data_only=True, read_only=True)
- otvara .xls sa xlrd
- cita SVE sheetove
- proizvodi DocumentEvidence sa cell Evidence listom

Legacy core/documents/excel_document.py ostaje netaknut do A7 (cleanup).
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import xlrd  # type: ignore[import-not-found]
from openpyxl import load_workbook  # type: ignore[import-not-found]

from parser_studio.domain.evidence import DocumentEvidence, Evidence, Locator
from parser_studio.ports.document_reader import DocumentReader

_XLSX_EXT = {".xlsx", ".xlsm"}
_XLS_EXT = {".xls"}
EXCEL_EXTENSIONS = _XLSX_EXT | _XLS_EXT


class ExcelDocumentReader:
    """Cita .xlsx/.xls fajlove i proizvodi DocumentEvidence."""

    reader_id: str = "excel"

    def __init__(self) -> None:
        self._cache: dict[str, DocumentEvidence] = {}

    def supports(self, path: Path) -> bool:
        """Prihvata .xlsx, .xlsm, .xls (case-insensitive sufiks)."""
        return path.suffix.lower() in EXCEL_EXTENSIONS

    def read(self, path: Path) -> DocumentEvidence:
        """Procitaj fajl i vrati DocumentEvidence; kesira po apsolutnoj putanji."""
        if not path.exists():
            raise FileNotFoundError(f"Fajl ne postoji: {path}")
        cache_key = str(path.resolve())
        if cache_key in self._cache:
            return self._cache[cache_key]
        suffix = path.suffix.lower()
        if suffix in _XLSX_EXT:
            doc = self._read_xlsx(path)
        elif suffix in _XLS_EXT:
            doc = self._read_xls(path)
        else:
            raise ValueError(f"Nepodržani format: {suffix}")
        self._cache[cache_key] = doc
        return doc

    def _read_xlsx(self, path: Path) -> DocumentEvidence:
        """Cita .xlsx sa openpyxl (data_only=True, read_only=True)."""
        wb = load_workbook(str(path), data_only=True, read_only=True)
        try:
            cells: list[Evidence] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        locator = Locator(
                            source_path=str(path),
                            kind="excel",
                            sheet=ws.title,
                            row=cell.row,
                            col=cell.column,
                        )
                        raw_text = "" if cell.value is None else str(cell.value).strip()
                        evidence = Evidence(
                            raw_value=cell.value,
                            raw_text=raw_text,
                            locator=locator,
                            element_type="excel_cell",
                            source_id=self.reader_id,
                        )
                        cells.append(evidence)
            return DocumentEvidence(
                document_id=str(path.resolve()),
                source_path=str(path),
                pages=[ws.title for ws in wb.worksheets],
                cells=cells,
            )
        finally:
            wb.close()

    def _read_xls(self, path: Path) -> DocumentEvidence:
        """Cita .xls sa xlrd (stariji binary format)."""
        wb = xlrd.open_workbook(str(path))
        cells: list[Evidence] = []
        sheet_names: list[str] = []
        try:
            for sheet in wb.sheets():
                sheet_names.append(sheet.name)
                for row_idx in range(sheet.nrows):
                    for col_idx in range(sheet.ncols):
                        value = sheet.cell_value(row_idx, col_idx)
                        if value == "" or value is None:
                            continue
                        locator = Locator(
                            source_path=str(path),
                            kind="excel",
                            sheet=sheet.name,
                            row=row_idx + 1,
                            col=col_idx + 1,
                        )
                        raw_text = str(value).strip() if value is not None else ""
                        evidence = Evidence(
                            raw_value=value,
                            raw_text=raw_text,
                            locator=locator,
                            element_type="excel_cell",
                            source_id=self.reader_id,
                        )
                        cells.append(evidence)
            return DocumentEvidence(
                document_id=str(path.resolve()),
                source_path=str(path),
                pages=sheet_names,
                cells=cells,
            )
        finally:
            pass  # xlrd nema close


__all__ = ["ExcelDocumentReader", "EXCEL_EXTENSIONS"]
