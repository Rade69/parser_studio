"""
Testovi za excel_headers engine.

Poseban naglasak na tri popravke u odnosu na original ExcelImporter.COLUMN_MAP:
normalizacija dijakritika, rangirano poklapanje umjesto golog substringa
(kratki aliasi ne smiju pogađati duže riječi), i jedna kolona = jedno polje.
"""
from __future__ import annotations

import pytest

from core.documents.excel_document import ExcelDocument
from core.engines.excel_headers import (
    ExcelHeadersEngine,
    _match_score,
    find_header_row,
    identify_columns,
    normalize_header,
)


class TestNormalizeHeader:
    def test_skida_dijakritike(self):
        assert normalize_header("Količina") == "kolicina"
        assert normalize_header("ŠIFRA") == "sifra"
        assert normalize_header("Težina") == "tezina"

    def test_skida_slovo_dj(self):
        assert normalize_header("Đubrivo") == "dubrivo"

    def test_skida_tacke(self):
        assert normalize_header("Kol.") == "kol"
        assert normalize_header("U.M.") == "um"
        assert normalize_header("Pos.") == "pos"

    def test_kolapsira_razmake(self):
        assert normalize_header("  Tarifni   broj  ") == "tarifni broj"

    def test_prazno_i_none(self):
        assert normalize_header(None) == ""
        assert normalize_header("") == ""
        assert normalize_header("   ") == ""


class TestMatchScore:
    def test_tacno_poklapanje_najvise(self):
        assert _match_score("kolicina", "kolicina") == 100.0

    def test_zasebna_rijec(self):
        assert _match_score("ukupna kolicina robe", "kolicina") == 90.0

    def test_pocetak_zaglavlja(self):
        assert _match_score("tarifni broj robe", "tarifni broj") == 95.0

    def test_kratak_alias_ne_pogadja_duzu_rijec(self):
        """KLJUČNO: 'kol' NE smije pogoditi 'kolona' (popravka #2 u odnosu na original)."""
        assert _match_score("kolona", "kol") == 0.0
        assert _match_score("vrijednost", "val") == 0.0
        assert _match_score("umjetnost", "um") == 0.0

    def test_kratak_alias_pogadja_tacno(self):
        assert _match_score("kol", "kol") == 100.0
        assert _match_score("um", "um") == 100.0

    def test_kratak_alias_pogadja_kao_zasebna_rijec(self):
        assert _match_score("neto kol robe", "kol") == 90.0

    def test_prazno_ne_pogadja(self):
        assert _match_score("", "kolicina") == 0.0
        assert _match_score("kolicina", "") == 0.0


class TestIdentifyColumns:
    def test_prepoznaje_osnovne_kolone(self):
        headers = ["Pos.", "Item Code", "Description", "U.M.", "Quantity", "Unit Price"]
        m = identify_columns(headers)
        assert m["product_code"] == 1
        assert m["naziv_robe"] == 2
        assert m["jm"] == 3
        assert m["kolicina"] == 4
        assert m["cijena_jed"] == 5

    def test_jedna_kolona_ne_moze_biti_dva_polja(self):
        m = identify_columns(["Naziv", "Opis"])
        assert len(set(m.values())) == len(m.values()), "kolone se ne smiju ponavljati"

    def test_jedno_polje_ne_moze_uzeti_dvije_kolone(self):
        m = identify_columns(["Cijena", "Unit Price"])
        assert list(m.values()).count(m.get("cijena_jed", -1)) == 1

    def test_srpski_nazivi_sa_dijakriticima(self):
        m = identify_columns(["Šifra", "Naziv robe", "Količina", "Zemlja porijekla"])
        assert m["product_code"] == 0
        assert m["naziv_robe"] == 1
        assert m["kolicina"] == 2
        assert m["zemlja_porijekla"] == 3

    def test_nepoznata_zaglavlja_se_ignorisu(self):
        m = identify_columns(["Nesto", "Bilo sta", "Xyz"])
        assert m == {}

    def test_prazne_celije_se_preskacu(self):
        m = identify_columns([None, "", "Naziv robe"])
        assert m == {"naziv_robe": 2}


@pytest.fixture
def sumaprom_like(tmp_path):
    """Struktura kao Šumaprom faktura — header podaci na vrhu, pa tabela."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.append(["SUMAPROM DOO BIJELJINA"])
    ws.append(["Invoice No.", "059/2022"])
    ws.append([])
    ws.append(["Pos.", "Item Code", "Description", "U.M.", "Quantity",
               "Unit Price", "Total Amount", "Country of origin"])
    ws.append([1, "A-100", "STARTNO UZE 4mm", "KOM", 5, 8.80, 44.00, "DE"])
    ws.append([2, "B-200", "NOZ KOSILICE MTD", "KOM", 10, 6.30, 63.00, "DE"])
    ws.append([3, "C-300", "SLAVINA GORIVA", "KOM", 10, 1.95, 19.50, "RS"])
    ws.append([])
    ws.append([None, None, "UKUPNO:", None, None, None, 126.50, None])

    path = tmp_path / "sumaprom.xlsx"
    wb.save(path)
    wb.close()
    return path


class TestFindHeaderRow:
    def test_nalazi_header_ispod_naslova(self, sumaprom_like):
        sheet = ExcelDocument(sumaprom_like).sheet(0)
        row, mapping = find_header_row(sheet)
        assert row == 3
        assert "naziv_robe" in mapping
        assert "kolicina" in mapping

    def test_vraca_minus_jedan_kad_nema_dovoljno_pogodaka(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.append(["Alfa", "Beta"])
        p = tmp_path / "bez_headera.xlsx"
        wb.save(p)
        wb.close()

        row, mapping = find_header_row(ExcelDocument(p).sheet(0))
        assert row == -1
        assert mapping == {}


class TestProbe:
    def test_prepoznaje_sumaprom_strukturu(self, sumaprom_like):
        result = ExcelHeadersEngine().probe(ExcelDocument(sumaprom_like))
        assert result.score > 0.5
        assert result.estimated_rows == 3
        assert result.suggested_config["header_row"] == 3
        assert "Invoice" in result.reason

    def test_prijedlog_konfiguracije_je_upotrebljiv_za_extract(self, sumaprom_like):
        doc = ExcelDocument(sumaprom_like)
        probe = ExcelHeadersEngine().probe(doc)
        table = ExcelHeadersEngine().extract(doc, probe.suggested_config)
        assert len(table) == probe.estimated_rows

    def test_prazan_dokument_daje_nulti_score(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.append(["nista", "ovdje"])
        p = tmp_path / "prazno.xlsx"
        wb.save(p)
        wb.close()

        result = ExcelHeadersEngine().probe(ExcelDocument(p))
        assert result.score == 0.0
        assert not result


class TestExtract:
    def test_izvlaci_samo_redove_sa_stavkama(self, sumaprom_like):
        doc = ExcelDocument(sumaprom_like)
        cfg = ExcelHeadersEngine().probe(doc).suggested_config
        table = ExcelHeadersEngine().extract(doc, cfg)

        assert len(table) == 3
        assert table.rows[0]["naziv_robe"].text == "STARTNO UZE 4mm"
        assert table.rows[2]["zemlja_porijekla"].text == "RS"

    def test_celije_nose_provenancu(self, sumaprom_like):
        doc = ExcelDocument(sumaprom_like)
        cfg = ExcelHeadersEngine().probe(doc).suggested_config
        table = ExcelHeadersEngine().extract(doc, cfg)

        cell = table.rows[0]["naziv_robe"]
        assert cell.row == 4
        assert cell.sheet == "Invoice"

    def test_odbaceni_redovi_imaju_razlog(self, sumaprom_like):
        doc = ExcelDocument(sumaprom_like)
        cfg = ExcelHeadersEngine().probe(doc).suggested_config
        table = ExcelHeadersEngine().extract(doc, cfg)

        assert table.rejected, "prazan red i UKUPNO red moraju biti odbaceni sa razlogom"
        razlozi = [r for _, r in table.rejected]
        assert any("prazan" in r for r in razlozi)

    def test_diagnostics_sadrzi_mapiranje(self, sumaprom_like):
        doc = ExcelDocument(sumaprom_like)
        cfg = ExcelHeadersEngine().probe(doc).suggested_config
        table = ExcelHeadersEngine().extract(doc, cfg)

        assert table.diagnostics["engine"] == "excel_headers"
        assert table.diagnostics["sheet"] == "Invoice"
        assert "naziv_robe" in table.diagnostics["columns"]

    @pytest.mark.parametrize(
        "footer_tekst",
        ["UKUPNO:", "Ukupno", "TOTAL", "Total Amount", "SVEGA", "PDV 17%",
         "Osnovica", "Za uplatu", "SUBTOTAL", "Napomena: roba je..."],
    )
    def test_footer_redovi_se_ne_racunaju_kao_stavke(self, tmp_path, footer_tekst):
        """Red zbira ima tekst u koloni naziva pa bi bez footer provjere bio tiho
        uvucen kao stavka — greska koja se u deklaraciji vidi tek kao pogresan zbir."""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["Sifra", "Naziv robe", "Kolicina", "Cijena"])
        ws.append(["A-1", "PRAVA STAVKA", 1, 10.0])
        ws.append([None, footer_tekst, None, 10.0])
        p = tmp_path / "sa_footerom.xlsx"
        wb.save(p)
        wb.close()

        doc = ExcelDocument(p)
        cfg = ExcelHeadersEngine().probe(doc).suggested_config
        table = ExcelHeadersEngine().extract(doc, cfg)

        assert len(table) == 1, f"'{footer_tekst}' je pogresno uvucen kao stavka"
        assert table.rows[0]["naziv_robe"].text == "PRAVA STAVKA"
        assert any("footera" in r for _, r in table.rejected)

    def test_stavka_koja_pocinje_slicno_footeru_se_ne_odbacuje(self, tmp_path):
        """'Totalizator' pocinje sa 'total' ali JESTE roba — ne smije se odbaciti
        samo zato sto dijeli prefiks sa footer pojmom."""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["Sifra", "Naziv robe", "Kolicina", "Cijena"])
        ws.append(["A-1", "TOTALIZATOR ZA KOSILICU", 1, 10.0])
        p = tmp_path / "slicno.xlsx"
        wb.save(p)
        wb.close()

        doc = ExcelDocument(p)
        cfg = ExcelHeadersEngine().probe(doc).suggested_config
        table = ExcelHeadersEngine().extract(doc, cfg)
        assert len(table) == 1

    def test_stop_after_empty_rows_prekida_na_kraju_tabele(self, sumaprom_like):
        doc = ExcelDocument(sumaprom_like)
        cfg = dict(ExcelHeadersEngine().probe(doc).suggested_config)
        cfg["stop_after_empty_rows"] = 1
        table = ExcelHeadersEngine().extract(doc, cfg)

        assert len(table) == 3
        assert any("kraj tabele" in r for _, r in table.rejected)
