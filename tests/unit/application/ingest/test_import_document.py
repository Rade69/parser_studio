# Tests: ingest/import_document.
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from parser_studio.adapters.documents.excel_reader import ExcelDocumentReader
from parser_studio.application.ingest.import_document import (
    ImportDocument,
    ImportRequest,
)


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Faktura"
    ws.append(["Naziv", "Kolicina"])
    ws.append(["Proizvod A", 5])
    path = tmp_path / "invoice.xlsx"
    wb.save(path)
    wb.close()
    return path


class TestImportDocumentConstruction:
    def test_requires_at_least_one_reader(self) -> None:
        with pytest.raises(ValueError, match="bar jedan DocumentReader"):
            ImportDocument(readers=())


class TestImportDocumentExecute:
    def test_execute_xlsx(self, sample_xlsx: Path) -> None:
        reader = ExcelDocumentReader()
        use_case = ImportDocument(readers=(reader,))
        request = ImportRequest(path=sample_xlsx)
        result = use_case.execute(request)
        assert result.reader_id == "excel"
        assert result.document.source_path == str(sample_xlsx)
        assert len(result.document.cells) > 0

    def test_execute_unsupported_format_raises(self, tmp_path: Path) -> None:
        reader = ExcelDocumentReader()
        use_case = ImportDocument(readers=(reader,))
        txt_file = tmp_path / "a.txt"
        txt_file.write_text("plain text")
        request = ImportRequest(path=txt_file)
        with pytest.raises(ValueError, match="ne podržava"):
            use_case.execute(request)

    def test_execute_uses_first_supporting_reader(self, sample_xlsx: Path) -> None:
        # Kreiraj drugi reader koji NE podržava xlsx (mock)
        from parser_studio.ports.document_reader import DocumentReader

        class NonSupportingReader:
            reader_id = "mock"
            def supports(self, path): return False
            def read(self, path): raise AssertionError("should not be called")

        excel_reader = ExcelDocumentReader()
        non_supporting = NonSupportingReader()
        use_case = ImportDocument(readers=(non_supporting, excel_reader))
        request = ImportRequest(path=sample_xlsx)
        result = use_case.execute(request)
        assert result.reader_id == "excel"

    def test_execute_document_id_in_request(self, sample_xlsx: Path) -> None:
        reader = ExcelDocumentReader()
        use_case = ImportDocument(readers=(reader,))
        request = ImportRequest(path=sample_xlsx, document_id="custom-id")
        result = use_case.execute(request)
        assert result.document.document_id == str(sample_xlsx.resolve())
