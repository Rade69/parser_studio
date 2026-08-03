"""
Testovi za ExcelDocument.

Fixture Excel fajlovi se prave U TESTU (openpyxl.Workbook), nikad se ne
commituju — stvarne fakture sadrže poslovne podatke (vidi .gitignore).
"""
from __future__ import annotations

import pytest

from core.documents.excel_document import EXCEL_EXTENSIONS, ExcelDocument


@pytest.fixture
def sample_xlsx(tmp_path):
    """Mali Excel koji imitira strukturu Šumaprom fakture: header podaci rasuti
    po vrhu, pa red sa nazivima kolona, pa stavke."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Faktura"
    ws.append(["SUMAPROM DOO", None, None, "Faktura br.", "059/2022"])
    ws.append([])
    ws.append(["Pos.", "Item Code", "Description", "U.M.", "Quantity", "Unit Price"])
    ws.append([1, "A-100", "STARTNO UZE 4mm", "KOM", 5, 8.80])
    ws.append([2, "B-200", "NOZ KOSILICE", "KOM", 10, 6.30])

    ws2 = wb.create_sheet("Prazan")
    ws2.append(["nista"])

    path = tmp_path / "sumaprom_uzorak.xlsx"
    wb.save(path)
    wb.close()
    return path


class TestOtvaranje:
    def test_can_open_prepoznaje_ekstenzije(self, tmp_path):
        assert ExcelDocument.can_open(tmp_path / "a.xlsx")
        assert ExcelDocument.can_open(tmp_path / "a.XLSX")
        assert ExcelDocument.can_open(tmp_path / "a.xls")
        assert not ExcelDocument.can_open(tmp_path / "a.pdf")

    def test_ekstenzije_pokrivaju_oba_formata(self):
        assert ".xlsx" in EXCEL_EXTENSIONS
        assert ".xls" in EXCEL_EXTENSIONS

    def test_ucitava_sve_sheetove(self, sample_xlsx):
        doc = ExcelDocument(sample_xlsx)
        assert doc.sheet_names == ["Faktura", "Prazan"]

    def test_sheet_po_indeksu_i_po_imenu(self, sample_xlsx):
        doc = ExcelDocument(sample_xlsx)
        assert doc.sheet(0).name == "Faktura"
        assert doc.sheet("Prazan").name == "Prazan"

    def test_nepostojeci_sheet_baca_keyerror(self, sample_xlsx):
        with pytest.raises(KeyError):
            ExcelDocument(sample_xlsx).sheet("NemaMe")

    def test_nepodrzana_ekstenzija_baca_valueerror(self, tmp_path):
        p = tmp_path / "a.txt"
        p.write_text("x")
        with pytest.raises(ValueError):
            ExcelDocument(p).sheets


class TestCitanjeCelija:
    def test_value_vraca_sirovu_vrijednost(self, sample_xlsx):
        s = ExcelDocument(sample_xlsx).sheet(0)
        assert s.value(0, 0) == "SUMAPROM DOO"
        assert s.value(3, 4) == 5

    def test_van_opsega_vraca_prazan_string_nikad_indexerror(self, sample_xlsx):
        s = ExcelDocument(sample_xlsx).sheet(0)
        assert s.value(999, 0) == ""
        assert s.value(0, 999) == ""
        assert s.value(-1, -1) == ""

    def test_none_postaje_prazan_string(self, sample_xlsx):
        s = ExcelDocument(sample_xlsx).sheet(0)
        assert s.value(0, 1) == ""

    def test_cell_nosi_provenancu(self, sample_xlsx):
        c = ExcelDocument(sample_xlsx).sheet(0).cell(3, 1)
        assert c.value == "A-100"
        assert (c.row, c.col, c.sheet) == (3, 1, "Faktura")
        assert c.text == "A-100"

    def test_row_text_spaja_neprazne_celije(self, sample_xlsx):
        s = ExcelDocument(sample_xlsx).sheet(0)
        assert s.row_text(0) == "SUMAPROM DOO Faktura br. 059/2022"
        assert s.row_text(1) == ""


class TestPretragaHeadera:
    def test_nalazi_red_sa_svim_pojmovima(self, sample_xlsx):
        s = ExcelDocument(sample_xlsx).sheet(0)
        assert s.find_row_containing(["Pos.", "Item Code", "Description"]) == 2

    def test_case_insensitive(self, sample_xlsx):
        s = ExcelDocument(sample_xlsx).sheet(0)
        assert s.find_row_containing(["pos.", "ITEM CODE"]) == 2

    def test_vraca_minus_jedan_kad_nema(self, sample_xlsx):
        s = ExcelDocument(sample_xlsx).sheet(0)
        assert s.find_row_containing(["Nepostojeca", "Kolona"]) == -1

    def test_require_all_false_trazi_bilo_koji(self, sample_xlsx):
        s = ExcelDocument(sample_xlsx).sheet(0)
        assert s.find_row_containing(["Nepostojeca", "SUMAPROM"], require_all=False) == 0

    def test_prazna_lista_pojmova_vraca_minus_jedan(self, sample_xlsx):
        assert ExcelDocument(sample_xlsx).sheet(0).find_row_containing([]) == -1


class TestKesiranje:
    def test_full_text_sadrzi_sve_sheetove(self, sample_xlsx):
        text = ExcelDocument(sample_xlsx).full_text()
        assert "SUMAPROM DOO" in text
        assert "STARTNO UZE" in text
        assert "Prazan" in text

    def test_ponovni_pristup_ne_cita_fajl_ponovo(self, sample_xlsx):
        doc = ExcelDocument(sample_xlsx)
        prvi = doc.sheets
        assert doc.sheets is prvi, "sheets mora biti kesiran, ne ucitan ponovo"

    def test_close_oslobadja_kes(self, sample_xlsx):
        doc = ExcelDocument(sample_xlsx)
        doc.full_text()
        doc.close()
        assert doc._sheets is None
        assert doc._full_text is None
        # nakon close mora ponovo raditi (lazy reload)
        assert doc.sheet_names == ["Faktura", "Prazan"]
